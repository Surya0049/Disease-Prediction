from __future__ import annotations

import csv
import os
import sqlite3
from datetime import datetime
from pathlib import Path
from typing import Any
from urllib.parse import urlencode

from flask import (
    Flask,
    flash,
    g,
    redirect,
    render_template,
    request,
    session,
    url_for,
)
from werkzeug.security import check_password_hash, generate_password_hash

BASE_DIR = Path(__file__).resolve().parent
DB_PATH = BASE_DIR / "data" / "app.db"

app = Flask(__name__)
app.secret_key = os.environ.get("FLASK_SECRET_KEY", "dev-secret-key")

DEFAULT_BACKGROUND = (
    "linear-gradient(135deg, rgba(22, 65, 148, 0.85), rgba(110, 67, 160, 0.85))"
)


def get_db() -> sqlite3.Connection:
    if "db" not in g:
        g.db = sqlite3.connect(DB_PATH)
        g.db.row_factory = sqlite3.Row
    return g.db


@app.teardown_appcontext
def close_db(exception: Exception | None) -> None:
    db = g.pop("db", None)
    if db is not None:
        db.close()


def init_db() -> None:
    DB_PATH.parent.mkdir(parents=True, exist_ok=True)
    db = sqlite3.connect(DB_PATH)
    db.execute(
        """
        CREATE TABLE IF NOT EXISTS users (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            username TEXT UNIQUE NOT NULL,
            password_hash TEXT NOT NULL,
            role TEXT NOT NULL
        )
        """
    )
    db.execute(
        """
        CREATE TABLE IF NOT EXISTS settings (
            id INTEGER PRIMARY KEY CHECK (id = 1),
            authority_name TEXT NOT NULL,
            background_css TEXT NOT NULL
        )
        """
    )
    db.execute(
        """
        CREATE TABLE IF NOT EXISTS dataset (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            disease TEXT NOT NULL,
            symptoms TEXT NOT NULL,
            medicine TEXT NOT NULL,
            dose TEXT NOT NULL,
            notes TEXT
        )
        """
    )
    admin_exists = db.execute("SELECT 1 FROM users WHERE username = ?", ("admin",)).fetchone()
    if not admin_exists:
        db.execute(
            "INSERT INTO users (username, password_hash, role) VALUES (?, ?, ?)",
            ("admin", generate_password_hash("admin123"), "admin"),
        )
    settings_exists = db.execute("SELECT 1 FROM settings WHERE id = 1").fetchone()
    if not settings_exists:
        db.execute(
            "INSERT INTO settings (id, authority_name, background_css) VALUES (1, ?, ?)",
            ("Health Authority", DEFAULT_BACKGROUND),
        )
    db.commit()
    db.close()


@app.before_request
def load_settings() -> None:
    init_db()
    db = get_db()
    settings = db.execute("SELECT * FROM settings WHERE id = 1").fetchone()
    g.settings = settings


def login_required(view):
    def wrapped_view(**kwargs):
        if "user_id" not in session:
            return redirect(url_for("login"))
        return view(**kwargs)

    wrapped_view.__name__ = view.__name__
    return wrapped_view


def admin_required(view):
    def wrapped_view(**kwargs):
        if "user_id" not in session:
            return redirect(url_for("login"))
        if session.get("role") != "admin":
            flash("Admin access required.", "error")
            return redirect(url_for("dashboard"))
        return view(**kwargs)

    wrapped_view.__name__ = view.__name__
    return wrapped_view


@app.route("/")
def index() -> Any:
    if "user_id" in session:
        return redirect(url_for("dashboard"))
    return redirect(url_for("login"))


@app.route("/login", methods=["GET", "POST"])
def login() -> Any:
    if request.method == "POST":
        username = request.form.get("username", "").strip()
        password = request.form.get("password", "").strip()
        db = get_db()
        user = db.execute("SELECT * FROM users WHERE username = ?", (username,)).fetchone()
        if user and check_password_hash(user["password_hash"], password):
            session.clear()
            session["user_id"] = user["id"]
            session["role"] = user["role"]
            session["username"] = user["username"]
            return redirect(url_for("dashboard"))
        flash("Invalid credentials. Please try again.", "error")
    return render_template("login.html")


@app.route("/logout")
def logout() -> Any:
    session.clear()
    return redirect(url_for("login"))


@app.route("/dashboard")
@login_required
def dashboard() -> Any:
    db = get_db()
    dataset_count = db.execute("SELECT COUNT(*) AS count FROM dataset").fetchone()["count"]
    user_count = db.execute("SELECT COUNT(*) AS count FROM users").fetchone()["count"]
    return render_template(
        "dashboard.html",
        dataset_count=dataset_count,
        user_count=user_count,
    )


@app.route("/predict", methods=["GET", "POST"])
@login_required
def predict() -> Any:
    db = get_db()
    dataset = db.execute("SELECT * FROM dataset ORDER BY disease").fetchall()
    symptoms_set = sorted(
        {symptom.strip() for row in dataset for symptom in row["symptoms"].split(",") if symptom}
    )
    prediction = None
    if request.method == "POST":
        age = request.form.get("age", "").strip()
        sex = request.form.get("sex", "").strip()
        selected_symptoms = request.form.getlist("symptoms")
        if not selected_symptoms:
            flash("Select at least one symptom for prediction.", "error")
        else:
            prediction = predict_disease(selected_symptoms, dataset)
            if prediction:
                query = f"{prediction['medicine']} {prediction['dose']}"
                search_url = f"https://www.google.com/search?{urlencode({'q': query})}"
                prediction["search_url"] = search_url
                prediction["age"] = age
                prediction["sex"] = sex
    return render_template(
        "predict.html",
        symptoms=symptoms_set,
        prediction=prediction,
    )


def predict_disease(selected_symptoms: list[str], dataset: list[sqlite3.Row]) -> dict[str, str] | None:
    if not dataset:
        return None
    best_match = None
    best_score = -1
    selected_set = {symptom.strip().lower() for symptom in selected_symptoms}
    for row in dataset:
        row_symptoms = {symptom.strip().lower() for symptom in row["symptoms"].split(",") if symptom}
        score = len(selected_set.intersection(row_symptoms))
        if score > best_score:
            best_score = score
            best_match = row
    if best_match is None or best_score == 0:
        return None
    return {
        "disease": best_match["disease"],
        "medicine": best_match["medicine"],
        "dose": best_match["dose"],
        "notes": best_match["notes"] or "Follow medical guidance for confirmation.",
    }


@app.route("/admin/users", methods=["GET", "POST"])
@admin_required
def admin_users() -> Any:
    db = get_db()
    if request.method == "POST":
        username = request.form.get("username", "").strip()
        password = request.form.get("password", "").strip()
        role = request.form.get("role", "user")
        if not username or not password:
            flash("Username and password are required.", "error")
        else:
            try:
                db.execute(
                    "INSERT INTO users (username, password_hash, role) VALUES (?, ?, ?)",
                    (username, generate_password_hash(password), role),
                )
                db.commit()
                flash("User created successfully.", "success")
            except sqlite3.IntegrityError:
                flash("Username already exists.", "error")
    users = db.execute("SELECT * FROM users ORDER BY username").fetchall()
    return render_template("admin_users.html", users=users)


@app.route("/admin/users/<int:user_id>/edit", methods=["GET", "POST"])
@admin_required
def edit_user(user_id: int) -> Any:
    db = get_db()
    user = db.execute("SELECT * FROM users WHERE id = ?", (user_id,)).fetchone()
    if user is None:
        flash("User not found.", "error")
        return redirect(url_for("admin_users"))
    if request.method == "POST":
        username = request.form.get("username", "").strip()
        role = request.form.get("role", "user")
        new_password = request.form.get("new_password", "").strip()
        if not username:
            flash("Username cannot be empty.", "error")
        else:
            db.execute("UPDATE users SET username = ?, role = ? WHERE id = ?", (username, role, user_id))
            if new_password:
                db.execute(
                    "UPDATE users SET password_hash = ? WHERE id = ?",
                    (generate_password_hash(new_password), user_id),
                )
            db.commit()
            flash("User updated.", "success")
            return redirect(url_for("admin_users"))
    return render_template("edit_user.html", user=user)


@app.route("/admin/users/<int:user_id>/delete", methods=["POST"])
@admin_required
def delete_user(user_id: int) -> Any:
    if session.get("user_id") == user_id:
        flash("You cannot delete your own account while logged in.", "error")
        return redirect(url_for("admin_users"))
    db = get_db()
    db.execute("DELETE FROM users WHERE id = ?", (user_id,))
    db.commit()
    flash("User deleted.", "success")
    return redirect(url_for("admin_users"))


@app.route("/admin/dataset", methods=["GET", "POST"])
@admin_required
def admin_dataset() -> Any:
    db = get_db()
    if request.method == "POST":
        if "dataset_file" in request.files:
            file = request.files["dataset_file"]
            if file.filename:
                added = import_dataset(file)
                if added:
                    flash(f"Imported {added} records.", "success")
                else:
                    flash("No records imported. Check file format.", "error")
            else:
                flash("Please choose a file to upload.", "error")
        else:
            disease = request.form.get("disease", "").strip()
            symptoms = request.form.get("symptoms", "").strip()
            medicine = request.form.get("medicine", "").strip()
            dose = request.form.get("dose", "").strip()
            notes = request.form.get("notes", "").strip()
            if not all([disease, symptoms, medicine, dose]):
                flash("Disease, symptoms, medicine, and dose are required.", "error")
            else:
                db.execute(
                    "INSERT INTO dataset (disease, symptoms, medicine, dose, notes) VALUES (?, ?, ?, ?, ?)",
                    (disease, symptoms, medicine, dose, notes),
                )
                db.commit()
                flash("Dataset record added.", "success")
    entries = db.execute("SELECT * FROM dataset ORDER BY disease").fetchall()
    return render_template("admin_dataset.html", entries=entries)


@app.route("/admin/dataset/<int:entry_id>/edit", methods=["GET", "POST"])
@admin_required
def edit_dataset(entry_id: int) -> Any:
    db = get_db()
    entry = db.execute("SELECT * FROM dataset WHERE id = ?", (entry_id,)).fetchone()
    if entry is None:
        flash("Dataset entry not found.", "error")
        return redirect(url_for("admin_dataset"))
    if request.method == "POST":
        disease = request.form.get("disease", "").strip()
        symptoms = request.form.get("symptoms", "").strip()
        medicine = request.form.get("medicine", "").strip()
        dose = request.form.get("dose", "").strip()
        notes = request.form.get("notes", "").strip()
        if not all([disease, symptoms, medicine, dose]):
            flash("Disease, symptoms, medicine, and dose are required.", "error")
        else:
            db.execute(
                "UPDATE dataset SET disease = ?, symptoms = ?, medicine = ?, dose = ?, notes = ? WHERE id = ?",
                (disease, symptoms, medicine, dose, notes, entry_id),
            )
            db.commit()
            flash("Dataset entry updated.", "success")
            return redirect(url_for("admin_dataset"))
    return render_template("edit_dataset.html", entry=entry)


@app.route("/admin/dataset/<int:entry_id>/delete", methods=["POST"])
@admin_required
def delete_dataset(entry_id: int) -> Any:
    db = get_db()
    db.execute("DELETE FROM dataset WHERE id = ?", (entry_id,))
    db.commit()
    flash("Dataset entry deleted.", "success")
    return redirect(url_for("admin_dataset"))


@app.route("/admin/settings", methods=["GET", "POST"])
@admin_required
def admin_settings() -> Any:
    db = get_db()
    if request.method == "POST":
        authority_name = request.form.get("authority_name", "").strip()
        background_css = request.form.get("background_css", "").strip()
        if not authority_name:
            flash("Authority name cannot be empty.", "error")
        else:
            db.execute(
                "UPDATE settings SET authority_name = ?, background_css = ? WHERE id = 1",
                (authority_name, background_css or DEFAULT_BACKGROUND),
            )
            db.commit()
            flash("Settings updated.", "success")
            return redirect(url_for("admin_settings"))
    return render_template("admin_settings.html", settings=g.settings, default_background=DEFAULT_BACKGROUND)


def import_dataset(file_storage) -> int:
    filename = file_storage.filename.lower()
    temp_path = BASE_DIR / "data" / f"upload_{datetime.utcnow().timestamp()}_{file_storage.filename}"
    file_storage.save(temp_path)
    added = 0
    db = get_db()
    if filename.endswith(".csv"):
        with temp_path.open(newline="", encoding="utf-8") as csvfile:
            reader = csv.DictReader(csvfile)
            for row in reader:
                if add_dataset_row(db, row):
                    added += 1
    elif filename.endswith(".xlsx") or filename.endswith(".xls"):
        if not is_openpyxl_available():
            flash("XLSX import requires openpyxl. Please install it first.", "error")
        else:
            import openpyxl

            workbook = openpyxl.load_workbook(temp_path)
            sheet = workbook.active
            headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
            for row in sheet.iter_rows(min_row=2, values_only=True):
                data = dict(zip(headers, row))
                if add_dataset_row(db, data):
                    added += 1
    else:
        flash("Unsupported file format. Use CSV or XLSX.", "error")
    db.commit()
    temp_path.unlink(missing_ok=True)
    return added


def is_openpyxl_available() -> bool:
    import importlib.util

    return importlib.util.find_spec("openpyxl") is not None


def add_dataset_row(db: sqlite3.Connection, row: dict[str, Any]) -> bool:
    disease = str(row.get("disease", "")).strip()
    symptoms = str(row.get("symptoms", "")).strip()
    medicine = str(row.get("medicine", "")).strip()
    dose = str(row.get("dose", "")).strip()
    notes = str(row.get("notes", "")).strip() if row.get("notes") is not None else ""
    if not all([disease, symptoms, medicine, dose]):
        return False
    db.execute(
        "INSERT INTO dataset (disease, symptoms, medicine, dose, notes) VALUES (?, ?, ?, ?, ?)",
        (disease, symptoms, medicine, dose, notes),
    )
    return True


if __name__ == "__main__":
    app.run(host="0.0.0.0", port=5000, debug=True)
